from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "documentacion"
ASSETS = OUT / "assets_generados"
LOGO = ROOT / "design" / "logo-mev-concept-v1.png"

SCREENSHOTS = {
    "mantenimiento": ROOT / "documentacion" / "assets_fuente" / "mantenimiento.jpg",
    "borradores": ROOT / "documentacion" / "assets_fuente" / "borradores.jpg",
    "huella": ROOT / "documentacion" / "assets_fuente" / "huella.jpg",
    "intervencion": ROOT / "documentacion" / "assets_fuente" / "intervencion.jpg",
    "revision": ROOT / "documentacion" / "assets_fuente" / "revision.jpg",
    "aprobados": ROOT / "documentacion" / "assets_fuente" / "aprobados.jpg",
}

NAVY = "17324D"
NAVY_2 = "244B6B"
ORANGE = "F28C28"
ORANGE_LIGHT = "FFF1E2"
BLUE_LIGHT = "EAF1F7"
GRAY_LIGHT = "F3F5F7"
GRAY = "5F6B76"
INK = "1C252C"
WHITE = "FFFFFF"
GREEN = "2E7D32"
RED = "B3261E"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_table_geometry(table, widths_dxa: Sequence[int], indent_dxa: int = 120) -> None:
    total = sum(widths_dxa)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, name="Calibri", size=None, color=INK, bold=None, italic=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def paragraph_keep_with_next(paragraph, keep=True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    node = p_pr.find(qn("w:keepNext"))
    if node is None:
        node = OxmlElement("w:keepNext")
        p_pr.append(node)
    node.set(qn("w:val"), "1" if keep else "0")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    set_run_font(run, size=9, color=GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)


def configure_styles(doc: Document, preset: str) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8 if preset == "thesis" else 6)
    normal.paragraph_format.line_spacing = 1.333 if preset == "thesis" else 1.25
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if preset == "thesis" else WD_ALIGN_PARAGRAPH.LEFT

    for name, size, before, after, color in (
        ("Title", 28, 0, 10, NAVY),
        ("Subtitle", 14, 0, 12, GRAY),
        ("Heading 1", 16, 18, 10, NAVY),
        ("Heading 2", 13, 12, 6, NAVY_2),
        ("Heading 3", 12, 8, 4, ORANGE),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name.startswith("Heading") or name == "Title"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Number"):
        style = styles[list_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375 if preset == "manual" else 0.5)
        style.paragraph_format.first_line_indent = Inches(-0.188 if preset == "manual" else -0.25)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25 if preset == "manual" else 1.208

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(GRAY)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(10)


def new_document(preset: str, running_title: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(doc, preset)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(running_title.upper())
    set_run_font(run, size=8.5, color=GRAY, bold=True)

    footer = section.footer
    add_page_number(footer.paragraphs[0])

    settings = doc.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)
    return doc


def add_cover(doc: Document, title: str, subtitle: str, meta: Sequence[str]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if LOGO.exists():
        shape = p.add_run().add_picture(str(LOGO), width=Inches(2.1))
        shape._inline.docPr.set("descr", "Identidad visual de MEV Mantenimiento")
        shape._inline.docPr.set("title", "Logotipo MEV")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(title)
    set_run_font(run, size=26, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(30)
    run = p.add_run(subtitle)
    set_run_font(run, size=14, color=NAVY_2)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("MAQUINARIAS · EQUIPOS · VEHÍCULOS")
    set_run_font(run, size=10, color=ORANGE, bold=True)
    p.paragraph_format.space_after = Pt(36)

    for line in meta:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(line)
        set_run_font(run, size=11, color=INK, bold=line.startswith("Luis Alberto"))
    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int = 1, page_break: bool = False) -> None:
    if page_break and len(doc.paragraphs) > 0:
        doc.add_page_break()
    p = doc.add_heading(text, level=level)
    paragraph_keep_with_next(p)


def add_paragraph(doc: Document, text: str, *, bold_prefix: str | None = None, italic=False,
                  align=None, after=None) -> None:
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r, italic=italic)


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_run_font(r)


def add_numbered(doc: Document, items: Iterable[str]) -> None:
    numbering = doc.part.numbering_part.element
    base_num_id = int(doc.styles["List Number"]._element.pPr.numPr.numId.val)
    base_num = next(node for node in numbering.findall(qn("w:num")) if int(node.get(qn("w:numId"))) == base_num_id)
    abstract_id = base_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    existing = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    new_num_id = max(existing, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_id)
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    for item in items:
        p = doc.add_paragraph(style="List Number")
        num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = 0
        num_pr.get_or_add_numId().val = new_num_id
        r = p.add_run(item)
        set_run_font(r)


def add_callout(doc: Document, title: str, text: str, kind="info") -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), ORANGE_LIGHT if kind == "warning" else BLUE_LIGHT)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), ORANGE if kind == "warning" else NAVY_2)
    borders.append(left)
    p_pr.append(borders)
    r = p.add_run(f"{title}: ")
    set_run_font(r, bold=True, color=NAVY)
    r = p.add_run(text)
    set_run_font(r)


def add_table(doc: Document, headers: Sequence[str], rows: Sequence[Sequence[str]], widths: Sequence[int],
              header_fill=BLUE_LIGHT) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=9.5, color=NAVY, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(str(value)) < 18 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, size=9.2)
        set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_toc(doc: Document, entries: Sequence[str]) -> None:
    add_heading(doc, "Contenido", 1)
    add_callout(doc, "Edición", "En Microsoft Word puede actualizar la tabla automática con clic derecho > Actualizar campo.")
    p = doc.add_paragraph()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), 'TOC \\o "1-3" \\h \\z \\u')
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "Tabla de contenido automática"
    r.append(t)
    fld.append(r)
    p._p.append(fld)
    doc.add_page_break()


def add_figure(doc: Document, path: Path, caption: str, width=3.15) -> None:
    if not path.exists():
        add_callout(doc, "Figura no disponible", caption, "warning")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", caption.split(".", 1)[0])
    c = doc.add_paragraph(style="Caption")
    c.add_run(caption)


def fit_font(size=32, bold=False):
    candidates = [
        Path(r"C:\Windows\Fonts\aptos.ttf"),
        Path(r"C:\Windows\Fonts\calibri.ttf"),
        Path(r"C:\Windows\Fonts\arial.ttf"),
    ]
    if bold:
        candidates = [
            Path(r"C:\Windows\Fonts\aptos-bold.ttf"),
            Path(r"C:\Windows\Fonts\calibrib.ttf"),
            Path(r"C:\Windows\Fonts\arialbd.ttf"),
        ] + candidates
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def centered(draw, box, text, font, fill=INK):
    x1, y1, x2, y2 = box
    bounds = draw.multiline_textbbox((0, 0), text, font=font, spacing=6, align="center")
    w, h = bounds[2] - bounds[0], bounds[3] - bounds[1]
    draw.multiline_text(((x1 + x2 - w) / 2, (y1 + y2 - h) / 2), text, font=font, fill="#" + fill,
                        spacing=6, align="center")


def draw_arrow(draw, start, end, color=NAVY_2, width=6):
    draw.line([start, end], fill="#" + color, width=width)
    x, y = end
    draw.polygon([(x, y), (x - 18, y - 10), (x - 18, y + 10)], fill="#" + color)


def make_diagrams() -> tuple[Path, Path]:
    ASSETS.mkdir(parents=True, exist_ok=True)
    architecture = ASSETS / "arquitectura_mev.png"
    workflow = ASSETS / "flujo_estados_mev.png"

    img = Image.new("RGB", (1600, 920), "white")
    d = ImageDraw.Draw(img)
    title = fit_font(46, True)
    box_font = fit_font(28, True)
    small = fit_font(24)
    d.text((70, 45), "Arquitectura funcional de MEV Mantenimiento", font=title, fill="#" + NAVY)
    boxes = [
        ((80, 180, 380, 340), "Usuario\nAndroid", ORANGE_LIGHT),
        ((500, 140, 1050, 380), "Interfaz Jetpack Compose\nPantallas y validaciones", BLUE_LIGHT),
        ((500, 500, 1050, 740), "Lógica de dominio\nParser de voz · Repositorios · Reportes", GRAY_LIGHT),
        ((1170, 180, 1520, 340), "Firebase\nAuthentication", BLUE_LIGHT),
        ((1170, 560, 1520, 720), "Cloud Firestore\nColecciones", ORANGE_LIGHT),
    ]
    for box, label, fill in boxes:
        d.rounded_rectangle(box, radius=26, fill="#" + fill, outline="#" + NAVY_2, width=4)
        centered(d, box, label, box_font if "\n" not in label else small)
    draw_arrow(d, (380, 260), (500, 260))
    draw_arrow(d, (775, 380), (775, 500))
    draw_arrow(d, (1050, 260), (1170, 260))
    draw_arrow(d, (1050, 620), (1170, 620))
    d.text((90, 820), "La voz se convierte en comandos; el usuario revisa los campos antes de guardar o enviar.",
           font=small, fill="#" + GRAY)
    img.save(architecture, quality=95)

    img = Image.new("RGB", (1600, 620), "white")
    d = ImageDraw.Draw(img)
    d.text((70, 45), "Ciclo de trazabilidad del registro", font=title, fill="#" + NAVY)
    states = [
        ((60, 230, 330, 390), "BORRADOR", GRAY_LIGHT),
        ((440, 230, 710, 390), "ENVIADO", BLUE_LIGHT),
        ((820, 120, 1090, 280), "APROBADO", "E9F5EA"),
        ((820, 360, 1090, 520), "DEVUELTO", "FCEBE9"),
        ((1200, 360, 1470, 520), "CORRECCIÓN", ORANGE_LIGHT),
    ]
    for box, label, fill in states:
        d.rounded_rectangle(box, radius=24, fill="#" + fill, outline="#" + NAVY_2, width=4)
        centered(d, box, label, box_font)
    draw_arrow(d, (330, 310), (440, 310))
    draw_arrow(d, (710, 280), (820, 210))
    draw_arrow(d, (710, 340), (820, 440))
    draw_arrow(d, (1090, 440), (1200, 440))
    d.line([(1335, 360), (1335, 310), (710, 310)], fill="#" + ORANGE, width=6)
    d.polygon([(710, 310), (730, 298), (730, 322)], fill="#" + ORANGE)
    img.save(workflow, quality=95)
    return architecture, workflow


def build_manual(architecture: Path, workflow: Path) -> Path:
    doc = new_document("manual", "MEV Mantenimiento | Manual de usuario")
    add_cover(
        doc,
        "Manual de usuario",
        "MEV Mantenimiento · Versión 1.1.0-beta06",
        [
            "Luis Alberto Barragan Villota",
            "Carrera de Desarrollo de Software",
            "Instituto Superior Tecnológico Rumiñahui",
            "Septiembre de 2026",
        ],
    )

    add_heading(doc, "Control del documento", 1)
    add_table(doc, ["Elemento", "Detalle"], [
        ["Producto", "MEV Mantenimiento"],
        ["Versión documentada", "1.1.0-beta06"],
        ["Plataforma", "Android 8.0 o superior"],
        ["Propósito", "Guía de instalación, operación y resolución de problemas"],
        ["Público", "Técnicos, vulcanizadores, supervisores, analistas y planificadores"],
        ["Estado", "Prototipo beta para evaluación académica"],
    ], [2100, 7260])
    add_callout(doc, "Importante", "La aplicación es una versión beta. No debe usarse como único repositorio de información crítica sin reglas de seguridad, respaldo y validación institucional.", "warning")
    add_toc(doc, [
        "1. Introducción", "2. Requisitos e instalación", "3. Acceso y roles", "4. Navegación general",
        "5. Mantenimiento", "6. Toma de huella", "7. Intervención de llanta", "8. Revisión y aprobados",
        "9. Reportes", "10. Comandos de voz", "11. Solución de problemas", "12. Buenas prácticas y glosario",
    ])

    add_heading(doc, "1. Introducción", 1)
    add_paragraph(doc, "MEV Mantenimiento es una aplicación móvil orientada al registro de actividades realizadas sobre maquinarias, equipos y vehículos. Integra mantenimiento mecánico y control de llantas en un mismo flujo de trabajo, con persistencia en la nube y trazabilidad por estado.")
    add_paragraph(doc, "El dictado es una ayuda de captura. La aplicación convierte la frase reconocida en comandos, actualiza campos y muestra un mensaje. La responsabilidad de revisar la información antes de guardarla o enviarla permanece en el usuario.")
    add_figure(doc, architecture, "Figura 1. Arquitectura funcional simplificada de la aplicación.", 6.2)

    add_heading(doc, "2. Requisitos e instalación", 1, page_break=True)
    add_heading(doc, "2.1 Requisitos", 2)
    add_bullets(doc, [
        "Teléfono o tableta con Android 8.0 (API 26) o superior.",
        "Conexión a Internet para autenticación y sincronización con Firebase.",
        "Servicio de reconocimiento de voz instalado y habilitado.",
        "Cuenta creada previamente en Firebase Authentication y perfil activo en Firestore.",
        "Permiso de micrófono cuando se utilice el dictado.",
    ])
    add_heading(doc, "2.2 Instalación del APK de demostración", 2)
    add_numbered(doc, [
        "Copie MEV-Mantenimiento-1.1.0-beta06-debug.apk al dispositivo.",
        "Abra el archivo desde el administrador de descargas.",
        "Si Android lo solicita, autorice temporalmente la instalación desde esa fuente.",
        "Pulse Instalar y espere la confirmación.",
        "Abra MEV Mantenimiento y conceda el permiso de micrófono al probar la voz.",
    ])
    add_callout(doc, "Seguridad", "El APK incluido es una compilación de depuración para demostración. Una distribución institucional debe usar una compilación release firmada.", "warning")

    add_heading(doc, "3. Acceso y roles", 1, page_break=True)
    add_heading(doc, "3.1 Inicio de sesión", 2)
    add_numbered(doc, [
        "Ingrese el correo registrado.",
        "Escriba la contraseña y pulse Iniciar sesión.",
        "La aplicación consulta el perfil asociado al UID del usuario.",
        "Si el perfil está activo, se abre el menú correspondiente al rol.",
    ])
    add_heading(doc, "3.2 Acceso por rol", 2)
    add_table(doc, ["Rol", "Funciones principales"], [
        ["Técnico mecánico", "Mantenimiento, borradores e historial"],
        ["Vulcanizador", "Huella, intervención, borradores e historial"],
        ["Supervisor", "Registro, revisión y aprobados"],
        ["Analista", "Revisión, aprobados y reportes"],
        ["Asistente", "Registro, vulcanización, revisión y reportes"],
        ["Planificador", "Flujos operativos, reportes y matriz base"],
        ["Jefe de operaciones", "Reportes y matriz base"],
        ["Gerente general", "Reportes y opciones gerenciales"],
    ], [2700, 6660])
    add_callout(doc, "Alcance beta", "Las opciones Usuarios y Panel gerencial aún no poseen un flujo completo conectado. No deben incluirse en la demostración principal.", "warning")

    add_heading(doc, "4. Navegación general", 1, page_break=True)
    add_bullets(doc, [
        "El encabezado identifica al usuario, cargo y rol.",
        "Cada tarjeta del menú usa un icono relacionado con la actividad.",
        "Los formularios se desplazan verticalmente; los botones finales pueden requerir bajar la pantalla.",
        "Volver al menú principal conserva la sesión; Cerrar sesión elimina la sesión local.",
        "Los mensajes superiores informan resultados, errores o comandos reconocidos.",
    ])
    add_figure(doc, workflow, "Figura 2. Estados y ciclo de corrección de un registro.", 6.2)

    add_heading(doc, "5. Registro de mantenimiento", 1, page_break=True)
    add_heading(doc, "5.1 Crear un registro", 2)
    add_numbered(doc, [
        "Abra Nuevo mantenimiento.",
        "Seleccione o dicte el activo.",
        "Elija Preventivo o Correctivo.",
        "Revise kilometraje y horómetro; complete el indicador que corresponda al activo.",
        "Registre el mantenimiento realizado, observaciones, orden de trabajo y número de pedido cuando aplique.",
        "Seleccione Guardar borrador si faltan datos o Enviar registro si está completo.",
    ])
    add_figure(doc, SCREENSHOTS["mantenimiento"], "Figura 3. Formulario de mantenimiento completado mediante comandos de voz durante una prueba en dispositivo.", 2.7)
    add_heading(doc, "5.2 Borradores y devoluciones", 2)
    add_paragraph(doc, "Mis borradores reúne mantenimientos, tomas de huella e intervenciones en estado BORRADOR o DEVUELTO. Un registro devuelto muestra el motivo indicado por el revisor y puede corregirse antes de enviarlo nuevamente.")
    add_figure(doc, SCREENSHOTS["borradores"], "Figura 4. Centro de borradores utilizado en las pruebas funcionales.", 2.7)
    add_callout(doc, "Consejo", "En la edición de borradores de mantenimiento puede volver a usar el micrófono para actualizar campos. Revise siempre el resultado antes del envío.")

    add_heading(doc, "6. Toma general de huella", 1, page_break=True)
    add_numbered(doc, [
        "Abra Toma general de huella.",
        "Seleccione el activo. La aplicación muestra únicamente las posiciones habilitadas para ese equipo.",
        "Complete proyecto, kilometraje u horómetro y las profundidades P1 a P12 disponibles.",
        "Ingrese estado general, novedad y nombre del técnico.",
        "Guarde como borrador o envíe a revisión.",
    ])
    add_callout(doc, "Validación", "Cada huella debe estar entre 0.1 y 26 mm. La interfaz bloquea valores mayores y el repositorio repite la validación al enviar.", "warning")
    add_figure(doc, SCREENSHOTS["huella"], "Figura 5. Selección de activo y posiciones dinámicas en una toma de huella.", 2.65)

    add_heading(doc, "7. Intervención de llanta", 1, page_break=True)
    add_paragraph(doc, "El módulo registra actividades individuales sobre una llanta. Los tipos disponibles son cambio, rotación, reparación, montaje, desmontaje y baja.")
    add_numbered(doc, [
        "Seleccione el activo y confirme su lectura operativa.",
        "Elija el tipo de intervención y la posición.",
        "Complete huella, marca, medida y serie cuando corresponda.",
        "Describa el motivo, observaciones y técnico responsable.",
        "Guarde como borrador o envíe a revisión.",
    ])
    add_figure(doc, SCREENSHOTS["intervencion"], "Figura 6. Formulario de intervención de llanta durante una prueba en dispositivo.", 2.65)

    add_heading(doc, "8. Revisión y registros aprobados", 1, page_break=True)
    add_heading(doc, "8.1 Revisar", 2)
    add_numbered(doc, [
        "Abra Revisión de registros.",
        "Verifique activo, lecturas, descripción y datos técnicos.",
        "Pulse Aprobar cuando la información sea correcta.",
        "Para devolver, escriba primero un motivo claro; el botón se habilita al existir texto.",
    ])
    add_figure(doc, SCREENSHOTS["revision"], "Figura 7. Revisión unificada de mantenimiento y huella.", 2.65)
    add_heading(doc, "8.2 Consultar aprobados", 2)
    add_paragraph(doc, "Registros aprobados consulta únicamente los documentos cuyo estado es APROBADO en mantenimiento, huellas e intervenciones. Esta separación evita confundir pendientes con información ya revisada.")
    add_figure(doc, SCREENSHOTS["aprobados"], "Figura 8. Toma de huella enviada en proceso de revisión y aprobación.", 2.65)

    add_heading(doc, "9. Reportes y exportación", 1, page_break=True)
    add_paragraph(doc, "Reportes consolida los tres tipos de registro, calcula indicadores operativos y permite filtrar por proyecto, activo, técnico, estado y rango de fechas.")
    add_bullets(doc, [
        "PDF: resumen general y páginas de detalle legibles para consulta.",
        "Excel: hojas de resumen, mantenimientos, huellas e intervenciones.",
        "Los archivos respetan los filtros aplicados al momento de exportar.",
        "La ubicación final la selecciona el usuario mediante el selector de archivos de Android.",
    ])
    add_callout(doc, "Interpretación", "Los contadores representan los documentos disponibles en Firestore. No equivalen por sí solos a ahorro, disponibilidad mecánica o productividad.")

    add_heading(doc, "10. Referencia de comandos de voz", 1, page_break=True)
    add_table(doc, ["Objetivo", "Ejemplo de frase", "Resultado esperado"], [
        ["Activo", "tractor 9", "MTRAC0009"],
        ["Activo", "mini 5", "MMINI0005"],
        ["Servicio", "servicio correctivo", "CORRECTIVO"],
        ["Trabajo", "mantenimiento realizado cambio de bandas", "Campo mantenimiento realizado"],
        ["Orden", "orden de trabajo OT-904", "Identificador en mayúsculas"],
        ["Pedido", "número de pedido 923", "923"],
        ["Huella", "posición 3 18 milímetros", "P3 = 18"],
        ["Intervención", "tipo de intervención rotación", "ROTACIÓN"],
        ["Acción", "guardar borrador", "Guarda sin enviar"],
        ["Acción", "enviar registro", "Envía a revisión"],
    ], [1900, 3800, 3660])
    add_callout(doc, "Dicción", "Use frases cortas, diga primero la etiqueta y luego el valor. El guion de una orden puede depender de lo que entregue el servicio de voz; confirme el campo en pantalla.", "warning")

    add_heading(doc, "11. Solución de problemas", 1, page_break=True)
    add_table(doc, ["Situación", "Acción recomendada"], [
        ["No inicia sesión", "Verifique Internet, correo, contraseña y que el perfil esté ACTIVO."],
        ["No aparece un activo", "Busque por código y confirme que exista en la colección activos."],
        ["El micrófono no inicia", "Conceda RECORD_AUDIO y confirme que exista un servicio de reconocimiento."],
        ["La voz interpreta otro valor", "Corrija manualmente, dicte una frase más corta y vuelva a comprobar."],
        ["No permite devolver", "Escriba el motivo de devolución antes de pulsar el botón."],
        ["Rechaza una huella", "Use un valor numérico entre 0.1 y 26 mm."],
        ["No exporta", "Conceda acceso al selector de archivos y elija una ubicación disponible."],
        ["No aparecen borradores", "Confirme que pertenezcan al usuario actual y estén BORRADOR o DEVUELTO."],
    ], [3100, 6260])

    add_heading(doc, "12. Buenas prácticas y glosario", 1, page_break=True)
    add_heading(doc, "12.1 Buenas prácticas", 2)
    add_bullets(doc, [
        "Verificar activo y lecturas antes de registrar una actividad.",
        "No compartir credenciales entre técnicos.",
        "Usar borrador cuando la información aún no esté completa.",
        "Indicar motivos de devolución concretos y accionables.",
        "Confirmar que un registro figure como aprobado antes de usarlo en un informe final.",
        "Evitar información sensible en observaciones de demostración.",
    ])
    add_heading(doc, "12.2 Glosario", 2)
    add_table(doc, ["Término", "Definición"], [
        ["Activo", "Maquinaria, equipo o vehículo identificado por un código único."],
        ["Huella", "Profundidad de la banda de rodamiento expresada en milímetros."],
        ["Borrador", "Registro editable que todavía no fue enviado a revisión."],
        ["Enviado", "Registro pendiente de decisión de un revisor."],
        ["Aprobado", "Registro revisado y aceptado."],
        ["Devuelto", "Registro que requiere corrección y contiene un motivo."],
        ["Parser", "Componente que transforma texto reconocido en comandos de formulario."],
    ], [2200, 7160])

    path = OUT / "Manual_de_Usuario_MEV_Mantenimiento.docx"
    doc.save(path)
    return path


def build_thesis(architecture: Path, workflow: Path) -> Path:
    doc = new_document("thesis", "MEV Mantenimiento | Trabajo de integración curricular")
    add_cover(
        doc,
        "Sistema de gestión de mantenimiento para flotas con generación automática de reportes mediante reconocimiento de voz",
        "Trabajo de integración curricular · Documento final",
        [
            "Autor: Luis Alberto Barragan Villota",
            "Carrera: Desarrollo de Software",
            "Tutor: Yngrid Melo",
            "Instituto Superior Tecnológico Rumiñahui",
            "Septiembre de 2026",
        ],
    )

    add_heading(doc, "Resumen", 1)
    add_paragraph(doc, "El presente trabajo desarrolla MEV Mantenimiento, una aplicación Android para registrar y dar seguimiento a actividades ejecutadas sobre maquinarias, equipos y vehículos. El problema abordado es la dispersión de registros operativos y la dificultad de capturar información estructurada en campo. La solución integra autenticación, menú por rol, catálogo de activos, mantenimiento preventivo y correctivo, toma de huella, intervención de llantas, borradores, revisión, historial, registros aprobados, reportes y exportación. Como aporte central se implementó un parser contextual que transforma frases reconocidas por el servicio de voz de Android en comandos verificables para los formularios.")
    add_paragraph(doc, "La metodología fue aplicada, tecnológica, iterativa e incremental, adaptando prácticas de Scrum al desarrollo individual. La validación combinó inspección funcional en dispositivo y pruebas unitarias. La compilación de la versión 1.1.0-beta06 se completó correctamente y se ejecutaron nueve pruebas automatizadas sin fallos, enfocadas en el reconocimiento de activos, el dictado compuesto y el rango permitido de huellas. Los resultados demuestran la viabilidad del flujo propuesto y la utilidad de mantener confirmación humana antes de guardar o enviar. No se reportan porcentajes de ahorro ni impacto productivo, porque todavía no se realizó una prueba controlada con operación real.")
    add_paragraph(doc, "Palabras clave: mantenimiento de flotas, Android, reconocimiento de voz, trazabilidad, Firebase, Jetpack Compose.", italic=True)

    add_heading(doc, "Abstract", 1, page_break=True)
    add_paragraph(doc, "This project develops MEV Mantenimiento, an Android application for recording and tracking maintenance activities performed on machinery, equipment, and vehicles. It addresses fragmented operational records and the difficulty of capturing structured information in the field. The solution combines authentication, role-based menus, asset catalogues, preventive and corrective maintenance, tire tread measurements, tire interventions, drafts, review, history, approved records, reports, and exports. Its main contribution is a contextual parser that converts Android speech-recognition results into verifiable form commands.")
    add_paragraph(doc, "An applied, technological, iterative, and incremental methodology was used, adapting Scrum practices to an individual academic project. Validation combined functional inspection on a physical device and unit tests. Version 1.1.0-beta06 compiled successfully and nine automated tests passed with no failures. The evidence supports the functional feasibility of the workflow and the value of human confirmation before saving or submitting data. Productivity or savings percentages are not claimed because no controlled real-operation study has yet been performed.")
    add_paragraph(doc, "Keywords: fleet maintenance, Android, speech recognition, traceability, Firebase, Jetpack Compose.", italic=True)

    add_toc(doc, [
        "Capítulo 1. Planteamiento del problema", "Capítulo 2. Marco conceptual y tecnológico",
        "Capítulo 3. Metodología y requisitos", "Capítulo 4. Diseño de la solución",
        "Capítulo 5. Implementación", "Capítulo 6. Pruebas, resultados y discusión",
        "Conclusiones", "Recomendaciones", "Referencias", "Anexos",
    ])

    add_heading(doc, "Capítulo 1. Planteamiento del problema", 1)
    add_heading(doc, "1.1 Contexto", 2)
    add_paragraph(doc, "La gestión de una flota mixta requiere conservar información sobre mantenimientos, lecturas de kilometraje u horómetro, condición de llantas, responsables y decisiones de revisión. Cuando estos datos se encuentran en formularios separados o se registran después de la jornada, aumentan las posibilidades de omisión, duplicidad y pérdida de trazabilidad.")
    add_paragraph(doc, "El trabajo en campo añade restricciones: el usuario puede llevar guantes, desplazarse entre equipos o disponer de poco tiempo para escribir. Por ello, la captura por voz resulta atractiva, pero no debe sustituir la verificación. Los servicios de reconocimiento pueden alterar palabras, números o signos; la interfaz debe convertir la voz en una propuesta editable y mantener al usuario como responsable de confirmar el dato.")
    add_heading(doc, "1.2 Formulación del problema", 2)
    add_paragraph(doc, "¿Cómo centralizar el registro y seguimiento de mantenimiento de una flota heterogénea, facilitando la captura de datos en campo mediante voz sin perder control, revisión y trazabilidad?")
    add_heading(doc, "1.3 Objetivo general", 2)
    add_paragraph(doc, "Desarrollar una aplicación Android para gestionar registros de mantenimiento y vulcanización de maquinarias, equipos y vehículos, incorporando reconocimiento de voz como mecanismo de captura asistida y reportes para consulta operativa.")
    add_heading(doc, "1.4 Objetivos específicos", 2)
    add_bullets(doc, [
        "Modelar activos, usuarios y registros de mantenimiento en una base documental.",
        "Implementar formularios de mantenimiento, toma de huella e intervención de llanta.",
        "Diseñar un parser contextual que transforme frases en valores de formulario.",
        "Establecer borradores, envío, aprobación y devolución para asegurar trazabilidad.",
        "Consolidar indicadores y exportar información filtrada a PDF y Excel.",
        "Validar las reglas críticas mediante pruebas automatizadas y pruebas en dispositivo.",
    ])
    add_heading(doc, "1.5 Justificación", 2)
    add_paragraph(doc, "La propuesta integra procesos que suelen tratarse por separado. El valor académico reside en combinar interfaz móvil declarativa, persistencia en la nube, reconocimiento de voz, reglas de negocio y control de estados en un producto ejecutable. El valor operativo esperado es disponer de información más estructurada y consultable; este beneficio se presenta como potencial y deberá medirse posteriormente en un piloto real.")
    add_heading(doc, "1.6 Alcance y limitaciones", 2)
    add_table(doc, ["Incluido en la beta", "Fuera del cierre actual"], [
        ["Inicio de sesión y menú por rol", "Recuperación de contraseña y alta completa de cuentas"],
        ["Mantenimiento y vulcanización", "Adjuntos fotográficos en almacenamiento seguro"],
        ["Borradores, revisión y aprobados", "Reglas de seguridad productivas validadas por rol"],
        ["Reportes y exportación", "Integración con ERP o sistema institucional"],
        ["Pruebas unitarias y en dispositivo", "Estudio controlado de tiempos y productividad"],
        ["Identidad visual e iconos", "Publicación firmada en tienda de aplicaciones"],
    ], [4680, 4680])

    add_heading(doc, "Capítulo 2. Marco conceptual y tecnológico", 1, page_break=True)
    add_heading(doc, "2.1 Gestión del mantenimiento", 2)
    add_paragraph(doc, "El mantenimiento preventivo agrupa actividades planificadas para reducir la probabilidad de fallas; el correctivo responde a una anomalía detectada. En ambos casos, la utilidad del registro depende de asociar el activo, la lectura operativa, el trabajo realizado, las observaciones y los documentos de referencia.")
    add_heading(doc, "2.2 Flota heterogénea", 2)
    add_paragraph(doc, "MEV significa maquinarias, equipos y vehículos. Esta clasificación exige evitar reglas rígidas: algunos activos se controlan por kilometraje, otros por horómetro y otros por ambos. De igual forma, la cantidad de posiciones de llanta varía por activo. El catálogo maestro es, por tanto, una fuente de configuración para los formularios.")
    add_heading(doc, "2.3 Trazabilidad", 2)
    add_paragraph(doc, "La trazabilidad se implementa mediante identificador de documento, UID del usuario, correo, marcas de tiempo y estado. BORRADOR permite completar; ENVIADO pone el registro a disposición del revisor; APROBADO confirma aceptación; DEVUELTO conserva el motivo y habilita corrección.")
    add_figure(doc, workflow, "Figura 1. Ciclo de estados utilizado por los registros.", 6.2)
    add_heading(doc, "2.4 Reconocimiento y procesamiento de voz", 2)
    add_paragraph(doc, "Android SpeechRecognizer proporciona acceso al servicio de reconocimiento del dispositivo y requiere permiso de micrófono. En MEV Mantenimiento el resultado textual pasa por un parser propio. La normalización elimina diferencias de acentuación y espacios, identifica etiquetas, extrae valores, convierte números expresados con palabras y produce comandos tipados. Esta separación permite probar el parser sin depender del micrófono.")
    add_heading(doc, "2.5 Tecnologías seleccionadas", 2)
    add_table(doc, ["Tecnología", "Uso en el proyecto", "Justificación"], [
        ["Kotlin", "Lenguaje principal", "Interoperabilidad con Android y sintaxis concisa"],
        ["Jetpack Compose", "Interfaz declarativa", "Componentes reutilizables, estado y Material 3"],
        ["Firebase Authentication", "Sesión", "Autenticación integrada con SDK Android"],
        ["Cloud Firestore", "Persistencia", "Modelo documental y consultas desde la app"],
        ["SpeechRecognizer", "Voz a texto", "API disponible en el entorno Android"],
        ["JUnit", "Pruebas", "Validación rápida de lógica determinista"],
    ], [2100, 3000, 4260])

    add_heading(doc, "Capítulo 3. Metodología y requisitos", 1, page_break=True)
    add_heading(doc, "3.1 Enfoque metodológico", 2)
    add_paragraph(doc, "El proyecto es aplicado y tecnológico: produce un artefacto de software para un problema operativo definido. El proceso fue iterativo e incremental, adaptando Scrum a un desarrollo individual. El Product Backlog se mantuvo mediante historias de usuario; cada incremento incluyó selección, diseño, implementación, validación y corrección.")
    add_paragraph(doc, "La definición de terminado exigió que el flujo pudiera ejecutarse, conservara los datos esperados, respetara la navegación por rol y no presentara errores críticos en el escenario probado. Una pantalla visible sin comportamiento conectado no se consideró funcionalidad terminada.")
    add_heading(doc, "3.2 Iteraciones", 2)
    add_table(doc, ["Iteración", "Objetivo", "Resultado al cierre"], [
        ["I1", "Problema, alcance y backlog", "Requisitos priorizados"],
        ["I2", "Arquitectura y base Android", "Proyecto ejecutable"],
        ["I3", "Acceso, perfiles y activos", "Sesión y catálogo"],
        ["I4", "Mantenimiento, borradores y voz", "Flujo operativo inicial"],
        ["I5", "Huella e intervención", "Vulcanización configurable"],
        ["I6", "Historial, reportes y exportación", "Consulta consolidada"],
        ["I7", "Correcciones, pruebas, diseño y documentación", "Versión beta06"],
    ], [1300, 4060, 4000])
    add_heading(doc, "3.3 Actores", 2)
    add_table(doc, ["Actor", "Necesidad principal"], [
        ["Técnico mecánico", "Registrar y corregir mantenimientos"],
        ["Vulcanizador", "Registrar huellas e intervenciones"],
        ["Supervisor", "Aprobar o devolver información"],
        ["Analista", "Consultar historial y generar reportes"],
        ["Planificador", "Configurar activos y supervisar el flujo"],
        ["Jefatura/Gerencia", "Consultar información consolidada"],
    ], [2600, 6760])
    add_heading(doc, "3.4 Requisitos funcionales", 2)
    add_table(doc, ["ID", "Requisito", "Estado beta06"], [
        ["RF-01", "Autenticar usuarios y recuperar su perfil", "Implementado"],
        ["RF-02", "Mostrar opciones según rol", "Implementado en interfaz"],
        ["RF-03", "Consultar catálogo de activos", "Implementado"],
        ["RF-04", "Registrar mantenimiento", "Implementado"],
        ["RF-05", "Registrar toma general de huella", "Implementado"],
        ["RF-06", "Registrar intervención de llanta", "Implementado"],
        ["RF-07", "Guardar y editar borradores", "Implementado"],
        ["RF-08", "Aprobar o devolver registros", "Implementado"],
        ["RF-09", "Consultar historial y aprobados", "Implementado"],
        ["RF-10", "Filtrar y exportar reportes", "Implementado"],
        ["RF-11", "Administrar usuarios desde la app", "Pendiente"],
        ["RF-12", "Adjuntar evidencia fotográfica", "Pendiente"],
    ], [900, 5600, 2860])
    add_heading(doc, "3.5 Requisitos no funcionales", 2)
    add_bullets(doc, [
        "Usabilidad: formularios legibles, desplazables y con mensajes de confirmación.",
        "Compatibilidad: Android 8.0 o superior.",
        "Integridad: validación de campos en interfaz y repositorio.",
        "Trazabilidad: usuario, estado y marcas temporales en los documentos.",
        "Mantenibilidad: separación entre pantallas, parser, repositorios y exportación.",
        "Seguridad: autenticación activa; reglas de Firestore productivas pendientes de endurecimiento.",
    ])

    add_heading(doc, "Capítulo 4. Diseño de la solución", 1, page_break=True)
    add_heading(doc, "4.1 Arquitectura", 2)
    add_paragraph(doc, "La aplicación usa una arquitectura por responsabilidades. Las pantallas Compose gestionan estado y eventos; el parser convierte texto en comandos; los repositorios construyen, validan y persisten documentos; Firebase aporta autenticación y almacenamiento.")
    add_figure(doc, architecture, "Figura 2. Arquitectura funcional de MEV Mantenimiento.", 6.2)
    add_heading(doc, "4.2 Colecciones de datos", 2)
    add_table(doc, ["Colección", "Contenido principal"], [
        ["usuarios", "Perfil, cargo, rol, estado y correo"],
        ["activos", "Código, tipo, subtipo, marca, modelo, indicador, lecturas y ubicación"],
        ["registros_mantenimiento", "Servicio, lecturas, trabajo, observaciones, orden, pedido y estado"],
        ["tomas_huella", "Proyecto, lecturas, P1-P12, estado general, novedad, técnico y estado"],
        ["intervenciones_llanta", "Tipo, posición, huella, marca, medida, serie, motivo, técnico y estado"],
    ], [3000, 6360])
    add_heading(doc, "4.3 Diseño del parser", 2)
    add_numbered(doc, [
        "Normalizar a minúsculas y retirar acentos.",
        "Detectar acciones globales como guardar o enviar.",
        "Identificar alias de activo y extraer su secuencia numérica.",
        "Localizar etiquetas de campo: proyecto, kilometraje, horómetro, servicio, orden, pedido, huella y técnico.",
        "Separar varias etiquetas dentro de una frase.",
        "Crear objetos VoiceCommand y aplicarlos en la pantalla activa.",
        "Mostrar confirmación y permitir corrección manual.",
    ])
    add_table(doc, ["Frase", "Comando"], [
        ["tractor 9", "SeleccionarActivo(MTRAC0009)"],
        ["mini 5", "SeleccionarActivo(MMINI0005)"],
        ["mantenimiento realizado cambio de bandas", "ActualizarAccionEjecutada(cambio de bandas)"],
        ["posición 3 18 milímetros", "ActualizarPosicion(3, 18)"],
        ["guardar borrador", "GuardarBorrador"],
    ], [4300, 5060])
    add_heading(doc, "4.4 Diseño visual", 2)
    add_paragraph(doc, "La identidad MEV utiliza azul oscuro para confianza y estructura, naranja para energía y trabajo operativo, y fondos de alto contraste. El símbolo integra referencias abstractas a maquinaria, equipos y vehículos sin utilizar logotipos institucionales. Los accesos incorporan iconos por actividad para favorecer reconocimiento visual.")

    add_heading(doc, "Capítulo 5. Implementación", 1, page_break=True)
    add_heading(doc, "5.1 Entorno de construcción", 2)
    add_table(doc, ["Parámetro", "Valor"], [
        ["Nombre de paquete", "com.luis.mevmantenimiento"],
        ["Versión", "1.1.0-beta06"],
        ["minSdk / targetSdk", "26 / 36"],
        ["compileSdk", "37"],
        ["Kotlin", "2.2.10"],
        ["Android Gradle Plugin", "9.2.1"],
        ["Firebase BoM", "34.15.0"],
        ["Interfaz", "Jetpack Compose + Material 3"],
    ], [3200, 6160])
    add_heading(doc, "5.2 Mantenimiento", 2)
    add_paragraph(doc, "El formulario integra activo, tipo de servicio, kilometraje, horómetro, mantenimiento realizado, observaciones, orden de trabajo y número de pedido. La lógica permite guardar como borrador o enviar; los registros devueltos se reabren con su motivo y pueden corregirse.")
    add_figure(doc, SCREENSHOTS["mantenimiento"], "Figura 3. Captura asistida en el módulo de mantenimiento.")
    add_heading(doc, "5.3 Vulcanización", 2)
    add_paragraph(doc, "La toma de huella consulta las posiciones habilitadas para el activo y conserva campos P1 a P12. La intervención registra una acción sobre una posición específica y puede incluir datos de la llanta. Ambos flujos comparten estados y revisión.")
    add_figure(doc, SCREENSHOTS["huella"], "Figura 4. Formulario de huella con posiciones configuradas por activo.")
    add_figure(doc, SCREENSHOTS["intervencion"], "Figura 5. Formulario de intervención de llanta.")
    add_heading(doc, "5.4 Revisión, historial y aprobados", 2)
    add_paragraph(doc, "La revisión unificada carga documentos ENVIADO de las tres colecciones. Aprobar cambia el estado y registra la fecha; devolver exige un motivo. La consulta de aprobados filtra por APROBADO y el historial del usuario mantiene enviados, aprobados y devueltos.")
    add_figure(doc, SCREENSHOTS["revision"], "Figura 6. Pantalla de revisión unificada.")
    add_heading(doc, "5.5 Reportes", 2)
    add_paragraph(doc, "El repositorio de reportes carga mantenimientos, huellas e intervenciones, aplica filtros y calcula indicadores. La exportación PDF produce resumen y detalle; la exportación Excel crea un libro con hojas separadas. Los valores exportados se derivan de los documentos filtrados.")
    add_heading(doc, "5.6 Validaciones", 2)
    add_bullets(doc, [
        "Sesión activa antes de escribir en Firestore.",
        "Activo obligatorio en los flujos enviados.",
        "Tipo de servicio o intervención cuando corresponde.",
        "Técnico y estado general obligatorios para huellas enviadas.",
        "Motivo obligatorio al devolver.",
        "Huella válida entre 0.1 y 26 mm en interfaz y repositorio.",
    ])

    add_heading(doc, "Capítulo 6. Pruebas, resultados y discusión", 1, page_break=True)
    add_heading(doc, "6.1 Estrategia de prueba", 2)
    add_paragraph(doc, "La validación combinó pruebas unitarias deterministas y pruebas funcionales en un dispositivo Android. Las unitarias se enfocaron en reglas que pueden verificarse sin red: interpretación de voz y rango de huellas. Las funcionales recorrieron formularios, borradores, envío, revisión y consulta.")
    add_heading(doc, "6.2 Resultado automatizado", 2)
    add_table(doc, ["Suite", "Casos", "Fallos", "Resultado"], [
        ["ExampleUnitTest", "1", "0", "Aprobado"],
        ["HuellaValidationTest", "4", "0", "Aprobado"],
        ["VoiceCommandParserTest", "4", "0", "Aprobado"],
        ["Total", "9", "0", "BUILD SUCCESSFUL"],
    ], [3800, 1500, 1500, 2560], header_fill="E9F5EA")
    add_paragraph(doc, "La ejecución testDebugUnitTest del 1 de septiembre de 2026 finalizó correctamente. Esta evidencia valida los casos codificados, pero no sustituye pruebas de integración de Firebase, rendimiento ni usabilidad.")
    add_heading(doc, "6.3 Matriz de pruebas funcionales", 2)
    add_table(doc, ["Caso", "Resultado esperado", "Evidencia"], [
        ["Tractor 9", "Código MTRAC0009", "Prueba unitaria aprobada"],
        ["Mini 5", "Código MMINI0005", "Prueba unitaria aprobada"],
        ["Mantenimiento realizado", "Completa el campo de trabajo", "Prueba unitaria y dispositivo"],
        ["Dictado compuesto", "Separa activo y mantenimiento", "Prueba unitaria aprobada"],
        ["Huella 26", "Aceptada", "Prueba unitaria aprobada"],
        ["Huella 27/40", "Rechazada", "Prueba unitaria y dispositivo"],
        ["Borrador", "Se conserva y puede editarse", "Prueba en dispositivo"],
        ["Aprobar/devolver", "Cambia estado; devolución exige motivo", "Prueba en dispositivo"],
        ["Aprobados", "Solo muestra APROBADO", "Prueba en dispositivo"],
    ], [2300, 4200, 2860])
    add_heading(doc, "6.4 Resultados funcionales", 2)
    add_bullets(doc, [
        "Se centralizaron tres tipos de registro en una aplicación Android.",
        "El parser reconoce alias de activos de flota mixta y construye códigos con secuencia de cuatro dígitos.",
        "El dictado puede separar varios campos cuando se expresan con etiquetas reconocibles.",
        "Los borradores de mantenimiento, huella e intervención se consultan en un centro unificado.",
        "La revisión distingue pendientes y aprobados, manteniendo el motivo de devolución.",
        "La validación de huellas se aplica de inmediato y vuelve a comprobarse antes de persistir.",
        "La identidad visual y los iconos mejoran la diferenciación de módulos.",
    ])
    add_figure(doc, SCREENSHOTS["borradores"], "Figura 7. Evidencia del centro de borradores en dispositivo.")
    add_figure(doc, SCREENSHOTS["aprobados"], "Figura 8. Evidencia del flujo de revisión de huellas.")
    add_heading(doc, "6.5 Discusión", 2)
    add_paragraph(doc, "El resultado respalda la viabilidad del enfoque, pero también confirma que la voz debe tratarse como captura asistida. Las pruebas detectaron errores concretos —alias de activos, mantenimiento realizado y límites de huella— que fueron convertidos en casos automatizados. Este ciclo representa inspección y adaptación, principios consistentes con Scrum.")
    add_paragraph(doc, "La arquitectura Firebase acelera el prototipo, aunque traslada una responsabilidad crítica a las reglas de seguridad. La visibilidad del menú por rol mejora la experiencia, pero no constituye autorización suficiente por sí sola. Antes de producción se debe impedir en backend que un usuario ejecute operaciones no permitidas.")
    add_paragraph(doc, "No se presentan reducciones porcentuales de tiempo, costos o fallas. Las capturas y pruebas demuestran comportamiento funcional; para afirmar impacto operativo se necesita una línea base, muestra de usuarios, periodo de observación y comparación controlada.")

    add_heading(doc, "Conclusiones", 1, page_break=True)
    add_numbered(doc, [
        "Se desarrolló una beta Android ejecutable que integra mantenimiento, vulcanización, revisión y reportes para una flota heterogénea.",
        "La separación entre reconocimiento de voz y parser contextual permitió transformar frases en comandos verificables y probar la lógica sin depender del micrófono.",
        "El flujo BORRADOR, ENVIADO, APROBADO y DEVUELTO aporta trazabilidad y habilita correcciones controladas.",
        "La configuración por activo resuelve diferencias de indicador y posiciones de llanta entre maquinarias, equipos y vehículos.",
        "Las nueve pruebas unitarias aprobadas y las pruebas en dispositivo aportan evidencia del comportamiento central de la versión beta06.",
        "La solución cumple el objetivo como prototipo académico; la seguridad productiva, la gestión completa de usuarios y la medición de impacto permanecen como pasos posteriores.",
    ])

    add_heading(doc, "Recomendaciones", 1, page_break=True)
    add_numbered(doc, [
        "Diseñar y desplegar reglas de Firestore que autoricen lecturas y escrituras por rol y propietario.",
        "Agregar recuperación de contraseña, alta administrada y auditoría de cambios de perfiles.",
        "Realizar un piloto con técnicos y vulcanizadores, midiendo tiempo de captura, correcciones y satisfacción.",
        "Ampliar la batería con pruebas instrumentadas de Compose, emuladores Firebase y escenarios sin conexión.",
        "Incorporar adjuntos fotográficos con validación de formato, tamaño y permisos.",
        "Generar una compilación release firmada y establecer política de versiones, respaldo y soporte.",
    ])

    add_heading(doc, "Referencias", 1, page_break=True)
    references = [
        "Android Developers. (2026). Develop UI with Jetpack Compose. https://developer.android.com/develop/ui",
        "Android Developers. (2026). Meet Android Studio. https://developer.android.com/studio/intro",
        "Android Developers. (2026). SpeechRecognizer. https://developer.android.com/reference/android/speech/SpeechRecognizer",
        "Barragan Villota, L. A. (2026a). Propuesta de tesis: Sistema de gestión de mantenimiento para flotas con generación automática de reportes mediante reconocimiento de voz [Documento académico no publicado]. Instituto Superior Tecnológico Rumiñahui.",
        "Barragan Villota, L. A. (2026b). Avance teórico N.° 4: MEV Mantenimiento [Documento académico no publicado]. Instituto Superior Tecnológico Rumiñahui.",
        "Firebase. (2026). Get started with Firebase Authentication on Android. https://firebase.google.com/docs/auth/android/start",
        "Firebase. (2026). Get started with Cloud Firestore. https://firebase.google.com/docs/firestore/quickstart",
        "JetBrains. (2025). Kotlin for Android. https://kotlinlang.org/docs/android-overview.html",
        "Pressman, R. S., & Maxim, B. R. (2020). Software engineering: A practitioner's approach (9th ed.). McGraw-Hill Education.",
        "Schwaber, K., & Sutherland, J. (2020). La Guía Scrum: La guía definitiva de Scrum. https://scrumguides.org/docs/scrumguide/v2020/2020-Scrum-Guide-Spanish-Latin-South-American.pdf",
    ]
    for ref in references:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(ref)
        set_run_font(r)

    add_heading(doc, "Anexo A. Ficha técnica", 1, page_break=True)
    add_table(doc, ["Elemento", "Valor"], [
        ["Aplicación", "MEV Mantenimiento"],
        ["Versión", "1.1.0-beta06"],
        ["Paquete", "com.luis.mevmantenimiento"],
        ["APK", "MEV-Mantenimiento-1.1.0-beta06-debug.apk"],
        ["Tamaño APK", "25,955,615 bytes"],
        ["SHA-256", "67448607B419202882EA7F13844D944D57AD72339368EB92F868011E3521810E"],
        ["Rama", "LuisB/registros-aprobados-validacion-huellas"],
        ["Último commit de identidad", "161b263 - Integra identidad visual e iconos MEV"],
    ], [3000, 6360])

    add_heading(doc, "Anexo B. Lista de comprobación de aceptación", 1, page_break=True)
    checks = [
        "La sesión válida abre un menú acorde al rol.",
        "El activo dictado coincide con el código esperado.",
        "El mantenimiento puede guardarse como borrador y enviarse.",
        "Los registros devueltos muestran el motivo y admiten corrección.",
        "Las posiciones de huella coinciden con el activo.",
        "Una huella mayor de 26 mm es rechazada.",
        "Una intervención exige tipo, posición y técnico al enviar.",
        "La devolución exige comentario.",
        "La consulta de aprobados excluye enviados y devueltos.",
        "Los filtros y exportaciones se ejecutan con datos consistentes.",
    ]
    for item in checks:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run("☐ " + item)
        set_run_font(r)

    add_heading(doc, "Anexo C. Guion resumido de demostración", 1, page_break=True)
    add_numbered(doc, [
        "Iniciar sesión y explicar el menú por rol.",
        "Crear mantenimiento con el dictado: tractor 9 mantenimiento realizado cambio de bandas.",
        "Guardar borrador, reabrir, corregir y enviar.",
        "Cambiar a rol revisor y aprobar o devolver.",
        "Consultar aprobados.",
        "Mostrar toma de huella y límite de 26 mm.",
        "Aplicar un filtro de reporte y enseñar exportación PDF/Excel.",
    ])
    add_callout(doc, "Plan alterno", "Si falla el servicio de voz, completar manualmente y continuar con un registro previamente enviado. La defensa debe demostrar el flujo, no depender de una única llamada de red.", "warning")

    add_heading(doc, "Anexo D. Declaración de alcance de evidencia", 1, page_break=True)
    add_paragraph(doc, "Las figuras de interfaz provienen de pruebas realizadas sobre la aplicación Android durante agosto de 2026. Los datos mostrados son de demostración. Las pruebas unitarias descritas corresponden a la ejecución local del 1 de septiembre de 2026. Este documento no atribuye resultados productivos a esos datos y distingue entre funcionalidad implementada, limitación y trabajo futuro.")

    path = OUT / "Tesis_Final_MEV_Mantenimiento.docx"
    doc.save(path)
    return path


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    architecture, workflow = make_diagrams()
    manual = build_manual(architecture, workflow)
    thesis = build_thesis(architecture, workflow)
    print(manual)
    print(thesis)


if __name__ == "__main__":
    main()
